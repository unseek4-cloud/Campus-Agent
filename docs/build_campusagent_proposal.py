from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "CampusAgent项目方案书.docx"

FONT_CN = "Microsoft YaHei"
FONT_BODY_CN = "SimSun"
FONT_LATIN = "Aptos"
COLOR_BLACK = "000000"
COLOR_NAVY = "17365D"
COLOR_BLUE = "2F5597"
COLOR_PALE_BLUE = "EAF1F8"
COLOR_PALE_GRAY = "F5F7FA"
COLOR_BORDER = "D9D9D9"
COLOR_MUTED = "666666"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 110, start: int = 110, bottom: int = 110, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = COLOR_BORDER, size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def set_table_widths(table, widths: Sequence[float]) -> None:
    total_twips = int(round(sum(widths) * 1440))
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid_columns = table._tbl.tblGrid.gridCol_lst
    for grid_col, width in zip(grid_columns, widths):
        grid_col.set(qn("w:w"), str(int(round(width * 1440))))

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(round(width * 1440))))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def apply_font_to_paragraph(paragraph, name: str = FONT_BODY_CN, size: float = 11, bold: bool | None = None, color: str = COLOR_BLACK) -> None:
    for run in paragraph.runs:
        set_run_font(run, name, size, bold, color)


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_field(run, field_name: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field_name
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])


def set_repeat_header_and_footer(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.text = "CampusAgent 项目方案书"
    apply_font_to_paragraph(paragraph, FONT_CN, 9, False, COLOR_MUTED)
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_run = paragraph.add_run()
    set_run_font(page_run, FONT_CN, 9, False, COLOR_MUTED)
    add_field(page_run, "PAGE")


def set_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.84)
    section.right_margin = Inches(0.84)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    set_repeat_header_and_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_BODY_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY_CN)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = FONT_CN
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    title_p_pr = title._element.get_or_add_pPr()
    title_p_bdr = title_p_pr.find(qn("w:pBdr"))
    if title_p_bdr is not None:
        title_p_pr.remove(title_p_bdr)

    for style_name, font_size, before, after in (
        ("Heading 1", 16, 14, 8),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11.5, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style.font.size = Pt(font_size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    list_style = styles["List Bullet"]
    list_style.font.name = FONT_BODY_CN
    list_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY_CN)
    list_style.font.size = Pt(11)
    list_style.paragraph_format.left_indent = Inches(0.28)
    list_style.paragraph_format.first_line_indent = Inches(-0.16)
    list_style.paragraph_format.space_after = Pt(3)

    caption = styles["Caption"]
    caption.font.name = FONT_CN
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    caption.font.size = Pt(9.5)
    caption.font.bold = True
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(COLOR_MUTED)
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_body(doc: Document, text: str, *, bold_lead: str | None = None, align=None, first_line: bool = True) -> None:
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    if first_line:
        paragraph.paragraph_format.first_line_indent = Inches(0.28)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, FONT_BODY_CN, 11, True, COLOR_BLACK)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, FONT_BODY_CN, 11, False, COLOR_BLACK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, FONT_BODY_CN, 11, False, COLOR_BLACK)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(item)
        set_run_font(run, FONT_BODY_CN, 11, False, COLOR_BLACK)


def add_numbered_steps(doc: Document, items: Sequence[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.3)
        paragraph.paragraph_format.first_line_indent = Inches(-0.3)
        paragraph.paragraph_format.keep_together = True
        marker = paragraph.add_run(f"{index}  ")
        set_run_font(marker, FONT_CN, 11, True, COLOR_BLUE)
        text = paragraph.add_run(item)
        set_run_font(text, FONT_BODY_CN, 11, False, COLOR_BLACK)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, FONT_CN, 9.5, True, COLOR_MUTED)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Sequence[float] | None = None,
    font_size: float = 9.5,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        set_table_widths(table, widths)
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for idx, header in enumerate(headers):
        cell = header_cells[idx]
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_margins(cell, 120, 105, 120, 105)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(header)
        set_run_font(run, FONT_CN, font_size, True, "FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        if widths:
            set_table_widths(table, widths)
        fill = "FFFFFF" if row_index % 2 == 0 else COLOR_PALE_BLUE
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            set_cell_shading(cell, fill)
            set_cell_margins(cell, 105, 105, 105, 105)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_run_font(run, FONT_BODY_CN, font_size, False, COLOR_BLACK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_layer_architecture(doc: Document) -> None:
    add_caption(doc, "图表 1  CampusAgent 总体架构分层")
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_table_widths(table, [1.45, 5.25])
    layers = [
        ("访问与交互层", "Vue 3  Vite  Element Plus  Axios  Router\n学生端交互  管理端入口  问答与办事流程展示"),
        ("业务与治理层", "Spring Boot 3 模块化单体\n身份权限  政策管理  事项管理  问答编排  审计追溯"),
        ("AI 能力层", "FastAPI 独立服务\n文档解析  OCR  切块  Embedding  检索  Rerank  RAG  LLM"),
        ("数据与索引层", "MySQL 保存业务事实与版本元数据\nChroma 或 FAISS 保存开发期向量索引  预留 Qdrant 适配器"),
        ("运行与交付层", "环境变量配置  统一日志  健康检查  测试  Docker 预留  CI 预留"),
    ]
    for index, (layer, content) in enumerate(layers):
        set_row_cant_split(table.rows[index])
        left, right = table.rows[index].cells
        left.width = Inches(1.45)
        right.width = Inches(5.25)
        left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(left, COLOR_NAVY if index in (0, 2) else COLOR_BLUE)
        set_cell_shading(right, "FFFFFF" if index % 2 == 0 else COLOR_PALE_BLUE)
        for cell in (left, right):
            set_cell_margins(cell, 145, 130, 145, 130)
        p = left.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(layer)
        set_run_font(run, FONT_CN, 10, True, "FFFFFF")
        p = right.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(content)
        set_run_font(run, FONT_BODY_CN, 9.8, False, COLOR_BLACK)
    doc.add_paragraph()


def add_directory_tree(doc: Document) -> None:
    lines = [
        "Campus-Agent/",
        "  frontend/                  Vue 3 前端应用",
        "    src/app/                 应用装配与路由",
        "    src/modules/             按业务领域组织页面和状态",
        "    src/shared/              HTTP 客户端与通用资源",
        "  backend/                   Spring Boot Maven 多模块",
        "    campus-app/              唯一启动与装配模块",
        "    campus-shared/           最小共享契约与异常基础",
        "    campus-identity/         用户 权限 认证",
        "    campus-policy/           政策 文档 版本",
        "    campus-affair/           事项 流程 材料清单",
        "    campus-question/         问答会话与反馈",
        "    campus-ai-integration/   AI 服务客户端与防腐层",
        "    campus-audit/            审计与追溯记录",
        "  ai-service/                FastAPI AI 服务",
        "    app/api/                 内部 REST 接口",
        "    app/core/                配置 异常 依赖装配",
        "    app/domain/              AI 领域对象",
        "    app/services/            处理流程和能力接口",
        "    app/infrastructure/      模型 OCR 解析器适配器",
        "    app/vector_store/        向量库抽象与实现",
        "    tests/                   单元与集成测试",
        "  docs/                      需求 架构 API ADR 测试文档",
        "  deploy/                    Docker 与部署配置",
        "  scripts/                   开发 运维 数据脚本",
    ]
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.12
        run = p.add_run(line)
        set_run_font(run, "Consolas", 8.6, False, COLOR_BLACK)


def add_section_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    remove_paragraph_borders(paragraph)
    for run in paragraph.runs:
        set_run_font(run, FONT_CN, {1: 16, 2: 13, 3: 11.5}[level], True, COLOR_BLACK)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(38)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026 第八届全球校园人工智能算法精英大赛")
    set_run_font(run, FONT_CN, 13, False, COLOR_BLACK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AI加软件创新方向")
    set_run_font(run, FONT_CN, 12, False, COLOR_MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = doc.styles["Title"]
    run = p.add_run("CampusAgent")
    set_run_font(run, FONT_CN, 29, True, COLOR_BLACK)
    remove_paragraph_borders(p)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("高校事务智能问答与办事引导系统项目方案书")
    set_run_font(run, FONT_CN, 20, True, COLOR_BLACK)
    remove_paragraph_borders(p)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    run = p.add_run("面向校园政策可信问答和事项办理引导")
    set_run_font(run, FONT_CN, 12, False, COLOR_MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(118)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("方案版本  V1.0")
    set_run_font(run, FONT_CN, 10.5, False, COLOR_BLACK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("编制日期  2026年9月")
    set_run_font(run, FONT_CN, 10.5, False, COLOR_BLACK)
    doc.add_page_break()


def build_document() -> None:
    doc = Document()
    set_document_defaults(doc)
    props = doc.core_properties
    props.title = "CampusAgent 高校事务智能问答与办事引导系统项目方案书"
    props.subject = "2026 第八届全球校园人工智能算法精英大赛 AI加软件创新方向项目方案"
    props.author = "CampusAgent 项目组"
    props.keywords = "CampusAgent 高校事务 RAG 智能问答 办事引导"
    props.comments = ""

    add_cover(doc)

    add_section_heading(doc, "方案摘要")
    add_body(
        doc,
        "CampusAgent 面向高校学生在政策查询和事务办理中的高频困难，建设一个由可信政策知识库驱动的智能问答与办事引导系统。系统以学校正式发布的政策、通知和办事指南为知识依据，向学生提供问题理解、政策检索、答案生成、办理步骤说明、材料清单生成和来源追溯。对管理人员，系统提供文档导入、政策版本维护、知识索引更新和审计查询能力。",
    )
    add_body(
        doc,
        "项目采用模块化单体与独立 AI 服务的组合架构。Vue 前端负责交互，Spring Boot 后端负责用户、权限、政策和事项等业务事实，FastAPI 服务负责文档解析、向量化、检索、重排和生成式问答。浏览器只访问 Java 后端，Java 通过内部 REST API 调用 Python 服务，由此保持权限边界、数据责任和技术演进路径清晰。",
    )
    add_body(
        doc,
        "第一阶段以可验证的最小闭环为目标：完成政策文档管理和版本控制，建立文档解析与索引链路，实现带引用的政策问答、政策检索、办事流程与材料清单展示，并提供必要的权限和审计能力。方案不在早期引入微服务、复杂工作流平台或多智能体编排，以避免基础数据和质量评估尚未稳定时扩大系统复杂度。",
    )
    add_caption(doc, "表 1  项目关键信息")
    add_table(
        doc,
        ["项目项", "方案内容"],
        [
            ["项目名称", "CampusAgent 高校事务智能问答与办事引导系统"],
            ["参赛方向", "2026 第八届全球校园人工智能算法精英大赛 AI加软件创新"],
            ["服务对象", "高校学生为主要用户  学校业务管理人员为知识维护用户"],
            ["核心价值", "降低政策查找成本  减少办理遗漏  提升答复一致性  保留完整依据"],
            ["核心理念", "可问  可查  可办  可追溯"],
            ["总体架构", "Vue 前端  Spring Boot 模块化单体  FastAPI 独立 AI 服务"],
            ["第一阶段结果", "形成从政策入库到引用式问答和事项引导的可演示闭环"],
        ],
        [1.45, 5.25],
    )

    add_section_heading(doc, "1 项目背景和问题定义")
    add_section_heading(doc, "1 1 项目背景", 2)
    add_body(
        doc,
        "高校政策与办事信息通常分散在门户网站、二级部门通知、公众号文章、PDF 文件和线下说明中。学生需要先判断主管部门，再理解政策适用条件，最后核对办理入口、时间限制和材料要求。信息虽然存在，但入口分散、文本较长、更新频繁，导致咨询重复、理解偏差和材料遗漏。",
    )
    add_body(
        doc,
        "通用大模型可以生成流畅回答，但如果缺少学校内部政策知识、版本信息和引用依据，容易出现内容过时、来源不明或补全不存在规则的问题。校园事务具有明确的制度约束，系统必须让答案与政策原文建立可检查的对应关系，并在证据不足时明确提示人工确认。",
    )
    add_section_heading(doc, "1 2 主要问题", 2)
    add_bullets(
        doc,
        [
            "政策入口分散  学生难以快速找到适用于本校和当前学期的正式文件",
            "政策文本复杂  条件 例外 时间节点和材料要求容易被遗漏",
            "版本更替频繁  搜索结果可能混合现行文件与历史文件",
            "咨询口径不一  不同渠道对同一事项可能给出不一致说明",
            "生成结果难核验  缺少文件名称 版本 页码和原文片段时无法追溯",
            "管理维护成本高  文档更新后需要同步更新搜索和问答知识",
        ],
    )
    add_section_heading(doc, "1 3 解决原则", 2)
    add_body(
        doc,
        "CampusAgent 以正式政策为事实来源，以结构化事项为办理依据，以检索增强生成作为问答方式。系统先检索和筛选证据，再由模型组织回答；回答同时返回政策名称、版本、页码或段落位置。对于没有充分证据的问题，系统返回不确定说明和下一步咨询渠道，不将模型推断作为正式政策。",
    )

    add_section_heading(doc, "2 产品定位和建设目标")
    add_section_heading(doc, "2 1 产品定位", 2)
    add_body(
        doc,
        "CampusAgent 是高校事务知识服务与办事引导平台，不替代学校职能部门作出审批决定，也不自动修改业务系统数据。它连接政策知识、事项规则和学生自然语言问题，帮助用户理解当前规则并完成办理准备。",
    )
    add_section_heading(doc, "2 2 建设目标", 2)
    add_table(
        doc,
        ["目标维度", "目标说明", "可验证结果"],
        [
            ["可信问答", "基于学校有效政策回答  证据不足时拒绝推断", "答案包含引用信息  可回到原文核验"],
            ["政策检索", "支持自然语言和关键词查找  兼顾语义相关性", "返回相关文件  版本和命中片段"],
            ["办事引导", "展示办理条件  步骤  材料  时间和责任部门", "用户可按清单逐项准备材料"],
            ["版本治理", "区分草稿 现行 失效和替代关系", "默认检索现行版本  历史答案可追溯"],
            ["模块演进", "业务模块和 AI 能力可独立替换", "更换模型或向量库不改写核心业务"],
        ],
        [1.05, 3.15, 2.5],
    )
    add_section_heading(doc, "2 3 非目标", 2)
    add_body(
        doc,
        "第一阶段不建设全校统一门户，不直接接管教务、财务和审批系统，不实现跨校知识共享，也不引入微服务治理平台。涉及个人学籍状态、审批结果或财务数据的查询，需要在后续取得明确的数据授权和接口条件后再进入范围。",
    )

    add_section_heading(doc, "3 用户角色和核心场景")
    add_caption(doc, "表 2  用户角色与任务")
    add_table(
        doc,
        ["角色", "主要任务", "系统支持"],
        [
            ["学生", "查询政策  理解条件  准备材料  确认办理入口", "自然语言问答  政策检索  事项卡片  材料清单  引用跳转"],
            ["知识管理员", "导入政策  维护元数据  发布新版本  处理失效文件", "文档处理  版本状态  索引更新  处理结果查看"],
            ["业务管理员", "维护事项条件  步骤  材料和责任部门", "事项配置  发布校验  变更记录"],
            ["审计人员", "核对回答依据  查询关键操作  分析争议问题", "请求记录  引用快照  版本链路  操作日志"],
            ["系统管理员", "管理账号  角色  配置和运行状态", "权限配置  健康检查  日志和告警入口"],
        ],
        [1.05, 2.55, 3.1],
    )
    add_section_heading(doc, "3 1 典型用户流程", 2)
    add_numbered_steps(
        doc,
        [
            "学生输入问题  例如奖学金申请条件或补办学生证所需材料",
            "Java 后端完成身份校验  参数校验和问题记录",
            "问答模块识别咨询类型并取得可用政策范围和事项数据",
            "AI 服务在指定知识库和有效版本中检索  必要时执行重排",
            "大模型仅根据选定证据生成答案并标注引用",
            "Java 后端合并结构化事项信息  权限可见范围和审计标识",
            "前端展示结论  条件  步骤  材料清单  来源和反馈入口",
        ],
    )

    add_section_heading(doc, "4 功能需求和阶段范围")
    add_section_heading(doc, "4 1 第一阶段 MVP", 2)
    add_table(
        doc,
        ["能力", "MVP 内容", "完成标准"],
        [
            ["账户与权限", "基础登录  JWT 角色校验  管理接口保护", "学生与管理员权限隔离  未授权请求被拒绝"],
            ["政策管理", "PDF 上传  元数据  发布状态  版本关系", "能识别现行与历史版本  更新后可重建索引"],
            ["文档处理", "文本提取  扫描页 OCR  清洗  切块  Embedding  入库", "文字型和扫描型 PDF 均可形成带页码索引"],
            ["政策检索", "关键词与向量检索  基础过滤  命中片段展示", "可按知识库和有效版本返回相关结果"],
            ["可信问答", "RAG 问答  引用  不足证据提示  简单反馈", "每个事实性回答都带来源或明确说明无法确认"],
            ["事项引导", "事项条件  步骤  材料清单  责任部门  办理入口", "用户可查看完整办理准备信息"],
            ["审计追溯", "记录问题  检索文档版本  引用  模型配置标识", "争议回答可重建当时使用的关键依据"],
            ["基础运维", "统一响应  异常处理  健康检查  配置外置", "三个服务可独立启动  故障可定位"],
        ],
        [1.05, 3.15, 2.5],
        9.1,
    )
    add_section_heading(doc, "4 2 后续扩展", 2)
    add_bullets(
        doc,
        [
            "多轮会话中的上下文澄清和事项槽位收集",
            "面向学校业务人员的政策差异对比和影响分析",
            "统一办事大厅或教务系统的授权跳转和进度查询",
            "知识质量看板  热点问题分析  无答案问题治理",
            "语音输入  图片材料识别和移动端适配",
            "Qdrant 集群  模型服务网关和更细粒度的多租户隔离",
        ],
    )
    add_body(
        doc,
        "这些能力均应建立在 MVP 的政策数据质量、问答评测结果和使用反馈之上。多智能体自动办理、自动提交审批和跨系统写入属于高风险扩展，不应在缺少授权和人工复核机制时提前实现。",
    )

    add_section_heading(doc, "5 总体技术架构")
    add_body(
        doc,
        "系统采用前后端分离、模块化单体和独立 AI 服务。该结构适合竞赛阶段和早期产品验证：部署单元数量有限，业务事务仍可由 Java 后端统一管理，同时 AI 依赖和 Python 运行环境不会进入 Java 进程。后续若访问量或团队边界发生明显变化，可依据监控数据拆分单个模块，而不是预先承担微服务复杂度。",
    )
    add_layer_architecture(doc)
    add_section_heading(doc, "5 1 关键架构约束", 2)
    add_bullets(
        doc,
        [
            "浏览器只调用 Java 业务 API  不直接访问 AI 服务",
            "MySQL 是用户 权限 政策元数据 事项和审计记录的事实来源",
            "向量数据库保存可重建索引  不承担业务状态和权限判断",
            "Python 不查询业务表  所需范围和标识由 Java 通过内部接口传入",
            "跨业务模块通过公开应用服务或契约协作  禁止访问其他模块的 Mapper Entity 和数据表",
            "模型 向量库 Embedding OCR 和文档解析器均通过接口适配",
            "配置和密钥由环境变量或部署配置注入  不进入源码和日志",
        ],
    )

    add_section_heading(doc, "6 前端架构设计")
    add_body(
        doc,
        "前端按业务模块组织页面、路由和接口调用，避免把所有页面和状态堆放在全局目录。app 目录只负责应用装配和路由，modules 目录承载身份、政策、事项、问答和审计等领域功能，shared 目录只保存稳定且跨模块复用的 HTTP 客户端、样式和通用组件。",
    )
    add_caption(doc, "表 3  前端目录职责")
    add_table(
        doc,
        ["目录", "职责", "依赖约束"],
        [
            ["src app", "应用入口  路由聚合  全局插件装配", "可以装配业务模块  不承载业务规则"],
            ["src modules auth", "登录  账号状态  权限路由", "只通过后端身份 API 获取状态"],
            ["src modules policy", "政策检索  文件详情  版本展示", "不直接调用 AI 内部接口"],
            ["src modules affair", "事项列表  办理步骤  材料清单", "通过业务 API 获取结构化数据"],
            ["src modules question", "问答会话  引用展示  反馈", "复用 shared HTTP  不读取其他模块私有状态"],
            ["src shared", "Axios 实例  通用组件  基础样式", "不能依赖任一具体业务模块"],
        ],
        [1.55, 2.65, 2.5],
    )
    add_section_heading(doc, "6 1 交互设计重点", 2)
    add_bullets(
        doc,
        [
            "答案先展示直接结论  再展示适用条件  办理步骤和材料",
            "引用信息可展开到文件名  版本  发布部门  页码和原文片段",
            "历史或失效政策使用明确状态标识  不与现行政策混排",
            "不确定回答说明缺少的证据以及建议咨询的部门",
            "耗时的文档处理任务显示处理中 成功和失败状态  避免重复提交",
        ],
    )

    add_section_heading(doc, "7 Java 业务后端设计")
    add_body(
        doc,
        "Java 后端是系统的业务边界和安全边界。campus-app 是唯一 Spring Boot 启动模块，负责配置和模块装配；其余 Maven 模块按领域划分，不单独成为应用。每个模块内部可采用 controller application domain infrastructure 的局部结构，但不建立全局 Controller Service Mapper 目录。",
    )
    add_caption(doc, "表 4  Java 模块职责和允许依赖")
    add_table(
        doc,
        ["模块", "核心职责", "允许依赖"],
        [
            ["campus-app", "应用启动  配置  模块装配  健康检查", "全部模块的公开装配入口"],
            ["campus-shared", "统一响应  通用异常  最小公共类型", "不依赖业务模块"],
            ["campus-identity", "用户 角色 权限 JWT 认证", "shared"],
            ["campus-policy", "政策文档 元数据 版本 生命周期", "shared  必要时依赖 identity 的公开鉴权契约"],
            ["campus-affair", "事项 条件 步骤 材料清单", "shared  可引用 policy 的公开文档标识"],
            ["campus-question", "问题 会话 反馈 问答编排", "shared  policy affair ai-integration 的公开应用接口"],
            ["campus-ai-integration", "AI REST 客户端 DTO 超时和错误映射", "shared  不依赖具体业务实现"],
            ["campus-audit", "关键操作和回答依据审计", "shared  通过事件或公开契约接收记录"],
        ],
        [1.52, 2.65, 2.53],
        9.0,
    )
    add_section_heading(doc, "7 1 禁止的直接依赖", 2)
    add_bullets(
        doc,
        [
            "任何模块不得直接访问其他模块的 Mapper Entity Repository 或数据库表",
            "campus-shared 不得成为业务工具集合  不放置政策 事项 问答等领域模型",
            "campus-question 不得绕过 campus-ai-integration 直接拼装 Python 接口",
            "campus-app 不编写业务规则  只完成启动 装配和少量基础设施配置",
            "Controller 不直接调用 Mapper  外部输入先经过应用服务和权限校验",
        ],
    )

    add_section_heading(doc, "8 Python AI 服务设计")
    add_body(
        doc,
        "Python 服务只负责 AI 能力，不维护用户、权限、事项状态或政策发布状态。其内部采用接口与适配器分离：services 定义文档提取、OCR、切块、Embedding、检索、重排和生成能力，infrastructure 封装第三方库，core 负责配置和依赖装配，api 只处理内部 REST 协议。",
    )
    add_section_heading(doc, "8 1 文档处理链路", 2)
    add_numbered_steps(
        doc,
        [
            "Java 为文档生成稳定标识并确认该版本允许建立索引",
            "AI 服务校验文件大小 类型 PDF 文件头和业务标识",
            "文字型页面优先提取文本  无有效文本的页面按需渲染并执行 OCR",
            "文本清洗保留页码和必要的段落边界  删除明显噪声",
            "切块策略在长度限制内优先选择段落和句子边界  保留适当重叠",
            "Embedding 适配器生成向量  每个块携带文档 版本 页码和哈希元数据",
            "向量存储按文档执行可重复替换  避免版本更新后产生重复索引",
        ],
    )
    add_section_heading(doc, "8 2 RAG 问答链路", 2)
    add_table(
        doc,
        ["阶段", "输入", "处理与输出"],
        [
            ["问题规范化", "用户问题  会话上下文", "提取主题  条件和时间信息  不改写用户关键约束"],
            ["检索过滤", "知识库  有效版本  权限范围", "关键词和向量召回  仅保留允许访问的文档"],
            ["重排", "候选片段", "按问题相关性和证据完整度重新排序"],
            ["上下文构建", "高相关片段  结构化事项", "去重  控制长度  保留来源标识"],
            ["答案生成", "受控提示词和证据", "生成结论 条件 步骤 材料和不确定说明"],
            ["引用校验", "答案句子和证据片段", "生成来源列表  过滤无法对应证据的事实性表达"],
            ["结果返回", "答案 引用  模型与检索标识", "返回 Java 进行审计落库和前端展示"],
        ],
        [1.25, 2.35, 3.1],
        9.2,
    )
    add_section_heading(doc, "8 3 可替换接口", 2)
    add_bullets(
        doc,
        [
            "EmbeddingProvider  隔离本地 Sentence Transformers 与外部向量模型服务",
            "VectorStore  隔离 Chroma FAISS 与后续 Qdrant",
            "Retriever 和 Reranker  允许独立评测召回与重排策略",
            "LLMProvider  隔离云端模型 本地模型和推理参数",
            "DocumentExtractor 和 OCRService  隔离 PDF 解析与 OCR 引擎",
        ],
    )

    add_section_heading(doc, "9 数据模型和知识治理")
    add_section_heading(doc, "9 1 业务数据边界", 2)
    add_table(
        doc,
        ["数据对象", "事实存储", "说明"],
        [
            ["用户 角色 权限", "MySQL", "由 identity 模块唯一管理"],
            ["政策文档和版本", "MySQL 与对象文件存储", "MySQL 保存元数据和状态  原始文件保存于受控文件存储"],
            ["事项 条件 步骤 材料", "MySQL", "结构化内容优先于从文本临时推断"],
            ["文本块和向量", "向量数据库", "可由原始文档和处理配置重建"],
            ["问答会话和反馈", "MySQL", "按隐私最小化原则保存"],
            ["回答证据快照", "MySQL 或审计存储", "保留文档版本 块标识 页码 模型和参数版本"],
        ],
        [1.55, 1.6, 3.55],
    )
    add_section_heading(doc, "9 2 政策版本规则", 2)
    add_bullets(
        doc,
        [
            "每份政策使用稳定 policyId  每次发布形成独立 versionId",
            "版本状态至少包括草稿 现行 失效  被替代  禁止物理覆盖已发布版本",
            "新版本生效时记录生效时间和替代关系  默认检索现行版本",
            "历史问答记录保存当时使用的 versionId  后续更新不改变原审计依据",
            "索引更新采用按文档替换和幂等处理  失败时保留上一份可用索引",
        ],
    )
    add_section_heading(doc, "9 3 可追溯字段", 2)
    add_body(
        doc,
        "关键回答建议记录 requestId、userId 的脱敏标识、questionId、knowledgeBaseId、policyId、versionId、documentId、chunkId、pageNumber、contentSha256、embeddingModel、rerankModel、llmModel、promptVersion、生成时间和反馈结果。该记录用于复核，不应保存与问题无关的个人信息或完整密钥。",
    )

    add_section_heading(doc, "10 Java 与 Python 通信边界")
    add_body(
        doc,
        "Java 与 Python 仅通过内部 REST API 通信。Java 负责决定谁可以调用、哪些知识可访问、哪个政策版本有效以及结果如何持久化；Python 负责在传入范围内完成 AI 计算。接口使用明确的版本前缀和统一错误结构，所有请求携带 requestId，便于跨服务定位。",
    )
    add_caption(doc, "表 5  内部 API 边界")
    add_table(
        doc,
        ["接口类别", "Java 提供", "Python 返回", "禁止事项"],
        [
            ["文档处理", "documentId knowledgeBaseId 文件 处理配置", "页数 块数 模型标识 哈希 处理状态", "Python 不修改政策发布状态"],
            ["检索", "问题 过滤条件 TopK 权限范围", "候选块 分数 来源元数据", "Python 不自行扩大知识范围"],
            ["问答", "问题 已授权上下文 生成策略", "答案 引用 不确定标识 模型标识", "Python 不写入用户会话表"],
            ["索引删除", "明确文档或版本标识 幂等键", "删除结果和受影响数量", "禁止模糊条件批量删除"],
            ["健康检查", "无或内部认证信息", "服务状态 版本 依赖就绪状态", "不返回密钥和内部路径"],
        ],
        [1.0, 2.2, 2.15, 1.35],
        8.8,
    )
    add_section_heading(doc, "10 1 接口可靠性约定", 2)
    add_bullets(
        doc,
        [
            "连接超时 读取超时和总处理超时分别配置  文档处理使用异步任务状态更合适",
            "查询类接口只在明确的瞬时错误下有限重试  写入类接口依赖幂等键",
            "错误码区分输入错误 文档错误 模型不可用 索引不可用和系统错误",
            "Java 将内部错误映射为对用户可理解的提示  不暴露堆栈和模型服务细节",
            "内部网络部署后使用服务身份或签名机制  不依赖前端传入的用户角色",
        ],
    )

    add_section_heading(doc, "11 安全 隐私和合规设计")
    add_section_heading(doc, "11 1 身份和权限", 2)
    add_body(
        doc,
        "Spring Security 和 JWT 用于身份认证和接口授权。权限判断以 Java 后端为准，前端路由控制只改善交互，不能作为安全措施。管理操作采用最小权限，政策发布、版本失效和索引删除等关键动作需要独立权限并写入审计日志。",
    )
    add_section_heading(doc, "11 2 数据安全", 2)
    add_bullets(
        doc,
        [
            "数据库密码 模型密钥和内部令牌通过环境变量或密钥管理系统注入",
            "上传文件校验扩展名 MIME 大小 PDF 文件头和解析结果  文件名不参与路径拼接",
            "日志对学号 手机号 邮箱和问题中的敏感信息进行脱敏或最小化保存",
            "外部模型调用前明确数据范围  不发送无关的用户身份和业务字段",
            "原始文档 业务数据库 向量索引和备份分别设置访问权限和保留周期",
        ],
    )
    add_section_heading(doc, "11 3 生成式 AI 风险控制", 2)
    add_bullets(
        doc,
        [
            "系统提示词明确限定证据来源  禁止根据常识补充校内规则",
            "检索结果为空或冲突时返回无法确认  并建议联系责任部门",
            "引用与答案分开校验  不能定位证据的事实性句子不得作为确定结论展示",
            "政策文件中的指令性文本按数据处理  防止文档内容改变系统指令",
            "对高影响事项保留人工确认提示  系统不承诺审批结果",
        ],
    )

    add_section_heading(doc, "12 非功能需求")
    add_table(
        doc,
        ["维度", "设计要求", "MVP 验证方式"],
        [
            ["可用性", "三个服务独立启动  核心依赖提供健康状态", "启动检查  健康接口  故障注入"],
            ["性能", "常规检索和问答在明确超时内完成  长任务异步化", "固定环境压测并记录 P50 P95"],
            ["可维护性", "按领域分模块  提供接口抽象  配置外置", "依赖检查  代码评审  模块测试"],
            ["可观测性", "统一 requestId 结构化日志  错误码  关键耗时", "跨服务日志能够关联一次请求"],
            ["可恢复性", "业务数据备份  向量索引可重建  索引替换幂等", "恢复演练和重建测试"],
            ["兼容性", "浏览器和 API 版本明确  文档处理兼容文字型和扫描型 PDF", "兼容性样本集测试"],
        ],
        [1.05, 3.05, 2.6],
    )
    add_body(
        doc,
        "性能目标需要在确定模型、硬件和网络条件后固化。方案阶段不把首次模型下载、首次 OCR 初始化等冷启动时间与稳定运行指标混在一起；竞赛演示环境应提前缓存模型并执行预热。",
    )

    add_section_heading(doc, "13 开发计划和里程碑")
    add_body(
        doc,
        "建议以八至十周完成第一阶段，按可运行闭环推进。每个阶段都交付可验证结果，避免先铺开全部页面和数据库表，再集中联调。实际日期可根据参赛节点和团队人数调整。",
    )
    add_caption(doc, "表 6  推荐实施计划")
    add_table(
        doc,
        ["阶段", "建议周期", "主要工作", "退出条件"],
        [
            ["阶段一 基线", "第1周", "冻结 MVP 术语 数据边界 接口规范 评测样本", "关键范围和非目标达成一致"],
            ["阶段二 身份和政策", "第2至3周", "登录权限 政策元数据 文件存储 版本生命周期", "可发布一份政策并追踪版本"],
            ["阶段三 知识入库", "第3至4周", "解析 OCR 清洗 切块 Embedding 向量入库", "文字型和扫描型 PDF 可形成索引"],
            ["阶段四 检索问答", "第5至6周", "混合检索 重排 RAG 引用 不确定处理", "基准问题可得到带来源回答"],
            ["阶段五 事项引导", "第6至7周", "事项数据 材料清单 问答与事项联动", "至少覆盖选定高频事项闭环"],
            ["阶段六 质量加固", "第8周", "权限 审计 异常 性能 安全 回归测试", "核心验收项通过且无高风险缺陷"],
            ["阶段七 竞赛交付", "第9至10周", "部署脚本 演示数据 使用说明 方案材料 演示流程", "可重复部署和完整演示"],
        ],
        [1.3, 1.0, 2.85, 1.55],
        8.7,
    )
    add_section_heading(doc, "13 1 推荐开发顺序", 2)
    add_numbered_steps(
        doc,
        [
            "先确定政策 版本 事项和引用的数据模型  同时建立少量真实评测问题",
            "完成身份权限和政策管理  确保 AI 输入来自可控版本",
            "稳定文档处理与向量入库  以可重复索引为验收重点",
            "分别评测召回 重排和生成  不把检索问题交给提示词掩盖",
            "实现事项引导和材料清单  与政策引用建立关联",
            "补齐审计 安全 性能和部署  最后完成演示脚本与材料",
        ],
    )

    add_section_heading(doc, "14 测试和评估方案")
    add_section_heading(doc, "14 1 测试分层", 2)
    add_bullets(
        doc,
        [
            "单元测试覆盖领域规则 切块策略 文本清洗 错误映射和权限判断",
            "模块集成测试覆盖 MySQL 映射 文件处理 向量存储和内部 API",
            "契约测试保证 Java DTO 与 Python Schema 在版本升级时兼容",
            "端到端测试覆盖政策发布 文档处理 问答 引用 事项清单和审计查询",
            "安全测试覆盖未授权访问 文件上传 边界输入 提示注入和敏感信息日志",
            "AI 评测使用固定问题集和人工标注证据  分开衡量检索与生成质量",
        ],
    )
    add_section_heading(doc, "14 2 建议验收指标", 2)
    add_table(
        doc,
        ["指标", "建议目标", "测量说明"],
        [
            ["可追溯覆盖率", "100%", "需要政策依据的正式回答均含引用或明确无法确认"],
            ["检索 Recall at K", "不低于 85%", "在人工标注验证集上  正确证据进入候选集合"],
            ["引用一致率", "不低于 90%", "引用片段能够支持对应事实性表达"],
            ["事项清单准确率", "不低于 95%", "与已发布结构化事项基准逐项比对"],
            ["无依据拒答率", "不低于 90%", "对知识库外或证据不足问题不生成确定校内规则"],
            ["关键接口成功率", "不低于 99%", "在约定测试负载和稳定依赖条件下统计"],
            ["高风险权限缺陷", "0 个", "未授权读取 管理操作绕过和跨模块数据越权"],
        ],
        [1.55, 1.35, 3.8],
        9.1,
    )
    add_body(
        doc,
        "以上数值是 MVP 的建议门槛，不代表现有系统已经达到。项目组应在样本规模、题型分布、模型版本和测试环境确定后冻结评测口径，并同时报告失败案例，避免只用平均分掩盖关键政策错误。",
    )

    add_section_heading(doc, "15 部署和运行方案")
    add_section_heading(doc, "15 1 开发环境", 2)
    add_body(
        doc,
        "开发阶段允许前端、Java 后端和 Python 服务分别在本机启动。MySQL 保存业务数据，Chroma 或 FAISS 保存向量索引。环境变量分别配置数据库连接、JWT 密钥、AI 服务地址、模型名称、向量库路径和文件大小限制。",
    )
    add_section_heading(doc, "15 2 交付环境", 2)
    add_body(
        doc,
        "交付阶段为三个服务分别准备容器构建文件，并通过 Docker Compose 或同等方式组织网络、配置、数据卷和健康检查。外部只暴露前端和 Java API，Python AI 服务位于内部网络。MySQL、文档文件和向量索引使用持久化卷并制定备份策略。",
    )
    add_section_heading(doc, "15 3 运行监控", 2)
    add_bullets(
        doc,
        [
            "服务健康 进程资源 数据库连接和向量库可用性",
            "请求总量 成功率 错误码 延迟和超时",
            "文档处理成功率 OCR 页数 切块数和索引耗时",
            "检索空结果比例 引用缺失比例 拒答比例和用户反馈",
            "模型令牌使用或推理耗时  便于评估成本和容量",
        ],
    )

    add_section_heading(doc, "16 项目组织和协作机制")
    add_table(
        doc,
        ["职责", "主要工作", "关键产出"],
        [
            ["产品与架构", "范围 数据口径 架构决策 接口边界 竞赛材料", "需求基线 ADR 方案书 演示流程"],
            ["前端开发", "学生端和管理端交互 引用展示 错误和加载状态", "Vue 应用 前端测试 使用说明"],
            ["Java 后端", "身份权限 领域模块 数据库 API AI 集成 审计", "Spring Boot 应用 数据迁移 OpenAPI"],
            ["AI 开发", "文档处理 检索 重排 RAG 模型适配 AI 评测", "FastAPI 服务 索引链路 评测报告"],
            ["测试与运维", "测试数据 自动化 部署 监控 安全检查", "测试报告 部署脚本 故障清单"],
        ],
        [1.2, 3.55, 1.95],
    )
    add_body(
        doc,
        "团队规模较小时可以一人承担多项职责，但验收口径仍需明确。影响接口和模块边界的变更通过 ADR 记录，API 以 OpenAPI 和契约测试为准，数据库变更使用版本化迁移脚本，避免依赖口头约定。",
    )

    add_section_heading(doc, "17 风险和应对措施")
    add_caption(doc, "表 7  主要风险清单")
    add_table(
        doc,
        ["风险", "原因和影响", "应对措施", "优先级"],
        [
            ["政策数据不完整", "知识库缺少正式文件或元数据  回答无法覆盖", "先选定高频事项和责任部门  建立文档清单与发布校验", "高"],
            ["版本混用", "旧文件仍被召回  可能给出过时规则", "强制版本状态过滤  保留替代关系和回答版本快照", "高"],
            ["模型幻觉", "证据不足时生成确定答案", "检索阈值 引用校验 拒答策略 人工评测", "高"],
            ["OCR 质量波动", "扫描件模糊 表格复杂 导致文本错误", "保留原页引用  低置信度标记  允许人工修订文本", "中"],
            ["AI 服务耗时", "模型加载和 OCR 占用较多资源", "延迟初始化 模型预热 任务异步化 限流和超时", "中"],
            ["模块边界侵蚀", "赶进度时跨模块访问数据", "公开契约 依赖检查 代码评审 ADR", "中"],
            ["敏感信息泄露", "问题或日志含个人信息  外部模型可能接收", "最小化传输 脱敏 权限 审计和部署策略", "高"],
            ["范围持续扩大", "过早接入多系统和自动办理", "以 MVP 验收项管理范围  扩展进入独立迭代", "中"],
        ],
        [1.2, 2.25, 2.55, 0.7],
        8.6,
    )

    add_section_heading(doc, "18 创新点和竞赛价值")
    add_section_heading(doc, "18 1 以版本化证据约束生成", 2)
    add_body(
        doc,
        "方案将 policyId、versionId、页码、内容哈希和模型配置纳入回答链路。检索默认面向现行版本，历史回答保留当时依据，使生成结果可以复核并解释政策更新带来的变化。创新点不在单纯接入大模型，而在把版本治理和回答证据放入同一数据链路。",
    )
    add_section_heading(doc, "18 2 结构化事项与政策文本协同", 2)
    add_body(
        doc,
        "材料清单、办理步骤和责任部门采用结构化数据维护，政策原文用于解释条件和提供引用。该设计减少模型从长文本中临时推断清单的风险，同时保留自然语言问答的易用性。",
    )
    add_section_heading(doc, "18 3 可替换的 AI 能力层", 2)
    add_body(
        doc,
        "Embedding、向量库、重排器、LLM、OCR 和解析器均通过接口隔离。竞赛阶段可以使用本地开源组件快速验证，后续根据数据规模和部署条件切换 Qdrant 或其他模型，而不改变政策和事项业务模块。",
    )
    add_section_heading(doc, "18 4 面向真实交付的工程闭环", 2)
    add_body(
        doc,
        "项目同时考虑权限、异常、审计、版本、评测和部署，不把演示限定为单个聊天页面。评审可以沿着政策导入、索引建立、问题回答、引用核验和版本更新观察完整闭环，并检查系统在证据不足时的处理方式。",
    )

    add_section_heading(doc, "19 交付物和验收条件")
    add_caption(doc, "表 8  第一阶段交付物")
    add_table(
        doc,
        ["交付物", "内容", "验收条件"],
        [
            ["前端应用", "学生问答与事项页面  管理基础页面", "可独立构建  核心流程可操作  错误状态明确"],
            ["Java 后端", "模块化单体 权限 政策 事项 问答编排 审计 API", "Java 21 构建通过  模块边界符合约束  API 统一"],
            ["AI 服务", "文档处理 检索 重排 RAG 模型和向量库适配", "FastAPI 可启动  核心接口测试通过  引用元数据完整"],
            ["数据库与样本", "迁移脚本 演示政策 事项和评测问题", "环境可重复初始化  样本来源和版本明确"],
            ["部署材料", "环境变量说明 容器配置 启动和备份说明", "新环境可按说明完成部署和健康检查"],
            ["质量材料", "测试报告 AI 评测报告 风险清单 已知限制", "指标口径可复现  失败案例和限制已记录"],
            ["竞赛材料", "项目方案书 演示脚本 系统说明和必要展示素材", "内容与实际实现一致  不夸大尚未完成能力"],
        ],
        [1.25, 3.15, 2.3],
        9.0,
    )

    add_section_heading(doc, "20 当前基础和下一步工作")
    add_body(
        doc,
        "当前仓库已经建立 frontend、backend、ai-service、docs、deploy 和 scripts 的基础结构。前端具备 Vue 3、Vite、Element Plus、Axios 和基础路由；Java 端已形成 Maven 多模块骨架、统一响应、全局异常和健康接口；AI 服务已具备 FastAPI 配置、统一 Schema、健康接口以及第一阶段 PDF 处理能力。",
    )
    add_body(
        doc,
        "AI 文档处理原型已经覆盖文字型 PDF 提取、扫描页按需 OCR、文本清洗、按边界切块、Embedding 和 Chroma 持久化，并通过服务接口与具体第三方库隔离。当前尚未形成完整业务闭环，检索、重排、LLM 问答、正式权限、政策版本数据库、事项管理和生产部署仍需按本方案推进。",
    )
    add_section_heading(doc, "20 1 近期优先事项", 2)
    add_numbered_steps(
        doc,
        [
            "冻结 MVP 数据字典和政策版本状态机  建立首批政策与事项样本",
            "实现身份权限和政策模块的数据库迁移  完成受控文档发布",
            "定义 Java 与 Python 的文档处理和检索契约  增加契约测试",
            "实现带元数据过滤的检索与重排  建立检索评测集",
            "实现引用式问答和无依据拒答  接入事项清单并完成端到端演示",
        ],
    )

    add_section_heading(doc, "结语")
    add_body(
        doc,
        "CampusAgent 的第一阶段重点是把高校正式政策转化为可检索、可引用、可维护的知识，并将问答结果与结构化办事信息连接起来。模块化单体能够控制业务复杂度，独立 AI 服务能够隔离模型和向量技术变化，两者通过稳定的内部 API 协作。只要优先保证政策版本、引用质量和评测闭环，项目即可形成可信且可持续扩展的竞赛成果。",
    )

    doc.add_page_break()
    add_section_heading(doc, "附录 A 项目目录结构")
    add_directory_tree(doc)

    add_section_heading(doc, "附录 B API 规范示例")
    add_section_heading(doc, "B 1 对外健康检查", 2)
    add_body(doc, "Java 健康检查使用 GET /api/v1/health。业务 API 统一返回 code、message、data、requestId 和 timestamp 等字段，错误响应保持同一外层结构。", first_line=False)
    add_section_heading(doc, "B 2 AI 内部健康检查", 2)
    add_body(doc, "AI 健康检查使用 GET /internal/v1/health。文档处理接口使用 POST /internal/v1/documents/process，仅允许 Java 后端或受控运维任务访问。", first_line=False)
    doc.add_page_break()
    add_section_heading(doc, "B 3 文档处理请求", 2)
    add_table(
        doc,
        ["字段", "类型", "说明"],
        [
            ["documentId", "string", "Java 管理的稳定文档标识"],
            ["knowledgeBaseId", "string", "所属知识库标识"],
            ["file", "PDF file", "待处理的政策文件"],
            ["requestId", "header string", "跨服务请求追踪标识"],
        ],
        [1.55, 1.35, 3.8],
    )
    add_body(
        doc,
        "正式开发时应以 OpenAPI 文件作为协议来源，示例只表达边界，不代替接口版本管理。新增字段优先保持向后兼容；破坏性变化使用新的 API 版本并明确迁移计划。",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
