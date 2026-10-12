"""Build editable administrative documents with consistent Chinese typography."""
import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from catalog import DATE, VERSION, POLICIES, SERVICES, TEMPLATES


def doc(title, department, purpose):
    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    for name, size in [('Normal', 11), ('Title', 18), ('Heading 1', 13), ('Heading 2', 11.5)]:
        style = d.styles[name]
        style.font.name = 'SimSun'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.space_before = Pt(0 if name in ('Normal', 'Title') else 6)
        style.paragraph_format.line_spacing = 1.15
        for border in list(style._element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
        if name != 'Normal':
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    d.styles['Title'].paragraph_format.space_after = Pt(12)
    p = d.add_paragraph(title, 'Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(f'归口部门 {department}    版本 {VERSION}    更新日期 {DATE}')
    p.paragraph_format.space_after = Pt(10)
    for r in p.runs:
        r.font.size = Pt(9)
    d.add_paragraph(purpose)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    for r in footer.runs:
        r.font.size = Pt(9)
    d.core_properties.author = '校园事务管理'
    d.core_properties.last_modified_by = '校园事务管理'
    d.core_properties.title = title
    d.core_properties.subject = purpose
    d.core_properties.keywords = '高校 学生事务 办理指南 材料'
    return d


def heading(d, title):
    d.add_paragraph(title, 'Heading 1')


def table(d, rows, widths, header=False):
    t = d.add_table(rows=0, cols=len(widths))
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Cm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + side)
        for k, v in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
            edge.set(qn('w:' + k), v)
        borders.append(edge)
    pr.append(borders)
    for index, values in enumerate(rows):
        row = t.add_row()
        cant = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant)
        if header and index == 0:
            repeat = OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                item = OxmlElement('w:' + side)
                item.set(qn('w:w'), '85')
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            tcpr.append(margins)
            if header and index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EDF3')
                tcpr.append(shade)
            cell.text = str(value)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.bold = bool(header and index == 0)
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def signatures(d, approvals=True):
    d.add_paragraph('本人承诺所填信息及提交材料真实完整，愿意配合办理所需的核验。')
    d.add_paragraph('申请人签名 __________________    日期 ______年____月____日')
    if approvals:
        table(d, [('受理与审核意见', '意见及经办人签名\n\n日期 ______年____月____日'),
                  ('办理结果或后续安排', '结果及办理人签名\n\n日期 ______年____月____日')], [4.0, 12.6])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=True)
    previous_path = root / '资料清单.json'
    previous = json.loads(previous_path.read_text(encoding='utf-8'))['documents'] if previous_path.exists() else []
    records = []

    def save(d, kind, key, title, folder):
        p = root / folder / (title + '.docx')
        p.parent.mkdir(parents=True, exist_ok=True)
        d.save(p)
        records.append(dict(kind=kind, key=key, title=title, docx=str(p.relative_to(root)),
                            pdf=str(p.with_suffix('.pdf').relative_to(root))))

    for p in POLICIES:
        d = doc(p['title'], p['department'], '本文件明确相关学生事务的办理范围、材料要求及受理职责，供学生准备申请和管理人员受理使用。具体项目日期、金额及学校专属要求以对应正式通知为准。')
        for title, body in p['sections']:
            heading(d, title)
            d.add_paragraph(body)
        heading(d, '配套事项与表单')
        d.add_paragraph('；'.join(s['name'] for s in SERVICES if s['policy'] == p['key']) + '。对应指南和表单可在事项详情中下载。')
        if p.get('reference'):
            ref = d.add_paragraph('参考依据 ' + ('普通高等学校学生管理规定 教育部令第41号' if p['key'] == 'academic' else '关于做好家庭经济困难学生认定工作的指导意见 教财〔2018〕16号') + '\n' + p['reference'])
            for r in ref.runs:
                r.font.size = Pt(8.5)
        save(d, 'policy', p['key'], p['title'], '01_管理制度')

    policy_by_key = {p['key']: p for p in POLICIES}
    for s in SERVICES:
        d = doc(s['guide_title'], s['department'], s['description'])
        heading(d, '办理条件与申请安排')
        for i, item in enumerate(s['requirements'], 1):
            d.add_paragraph(f'{i}. {item}')
        d.add_paragraph('申请时间 ' + s['deadline'] + '。')
        d.add_paragraph('办理入口 校园事务系统对应事项；线下受理由归口部门告知具体窗口。联系方式与工作时间请查看受理部门公开渠道。')
        heading(d, '所需材料')
        for i, m in enumerate(s['materials'], 1):
            d.add_paragraph(f"{i}. {m['name']} {'必交' if m['required'] else '按适用情形提交'}\n{m['description']}")
        if s['notes']:
            heading(d, '办理提醒')
            d.add_paragraph(s['notes'])
        d.add_page_break()
        heading(d, '办理流程')
        for i, step in enumerate(s['steps'], 1):
            d.add_paragraph(f"{i} {step['title']}", 'Heading 2')
            d.add_paragraph(step['description'] + ' 负责人员或单位 ' + step['department'] + '。')
        heading(d, '结果与后续操作')
        d.add_paragraph('办理完成后，请核对' + s['outcome'] + '。处理时间以受理反馈为准；材料补正期间应及时提交所缺内容。如需修改或撤回，先查看系统状态并联系受理部门。')
        d.add_paragraph('政策依据 ' + policy_by_key[s['policy']]['title'] + '。配套申请表及材料模板可从本事项材料清单下载。')
        save(d, 'guide', s['code'], s['guide_title'], '02_办理指南')

        d = doc(s['form_title'], s['department'], '请按本事项办理条件填写；无对应情形的填写不适用，审批意见由受理部门填写。')
        basic = [('姓名与学号', '姓名 __________________    学号 __________________'),
                 ('学院专业与班级', '________________________________________________'),
                 ('联系办法与申请日期', '联系办法 __________________    日期 ______________')]
        fields = [(f, '________________________________________________') for f in s['fields']]
        table(d, basic + fields, [4.6, 12.0])
        d.add_paragraph('随附材料名称 ____________________________________________________')
        signatures(d)
        save(d, 'form', s['code'], s['form_title'], '03_申请表')

    for item in TEMPLATES:
        d = doc(item['title'], '对应事项受理部门', item['purpose'])
        table(d, [(f, '________________________________________________') for f in item['fields']], [5.0, 11.6])
        if item.get('table'):
            count = len(item['table'])
            widths = [1.2, 4.2, 6.1, 1.4, 3.7] if count == 5 else [3.2, 4.2, 5.7, 3.5]
            table(d, [item['table']] + [['']*count for _ in range(6)], widths, header=True)
        d.add_paragraph(item['note'])
        if item['key'] == 'authorization':
            d.add_paragraph('委托人签名 __________________    受托人签名 __________________')
            d.add_paragraph('签署日期 ______年____月____日')
        else:
            signatures(d, approvals=False)
        save(d, 'template', item['key'], item['title'], '04_材料模板')

    d = doc('校园事务制度与办事资料目录', '校园事务管理', '本目录覆盖学籍、教务、奖助、请假、住宿、证件、心理支持、活动、借阅、就业和实验室事务。各事项配套管理制度、办理指南、申请表和适用材料模板。')
    heading(d, '资料使用')
    d.add_paragraph('先阅读制度和办理指南，再填写 Word 申请表。PDF 版本用于浏览与打印；Word 版本用于填写和维护。申请中只提交办理所需的资料，期限和项目参数按当期通知核对。')
    heading(d, '事项与政策对应目录')
    rows = [('事项名称', '归口部门', '对应管理文件')]
    rows.extend((s['name'], s['department'], policy_by_key[s['policy']]['title']) for s in SERVICES)
    table(d, rows, [4.2, 5.3, 7.1], header=True)
    heading(d, '材料模板')
    for t in TEMPLATES:
        d.add_paragraph(t['title'] + '。' + t['purpose'])
    heading(d, '维护要求')
    d.add_paragraph('事项名称、条件、材料与流程应与配套制度同步维护。政策更新后先核对受影响事项及表单，再更新来源文档和下载附件。学生个人申请资料与公开知识文件分开管理，申请资料不进入公共政策知识库。')
    save(d, 'catalog', 'catalog', '校园事务制度与办事资料目录', '00_资料目录')

    for r in records:
        p = root / r['docx']
        r['docx_sha256'] = hashlib.sha256(p.read_bytes()).hexdigest()
    manifest = dict(version=VERSION, updated_at=DATE, document_count=len(records),
                    service_count=len(SERVICES), policy_count=len(POLICIES), documents=records)
    (root / '资料清单.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    current_paths = {r[format] for r in records for format in ('docx', 'pdf')}
    for old in previous:
        for format in ('docx', 'pdf'):
            if old[format] not in current_paths:
                obsolete = (root / old[format]).resolve()
                if not obsolete.is_relative_to(root):
                    raise ValueError('Refusing to remove an artifact outside the output directory')
                obsolete.unlink(missing_ok=True)
    print(json.dumps({'documents': len(records), 'services': len(SERVICES), 'policies': len(POLICIES)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
